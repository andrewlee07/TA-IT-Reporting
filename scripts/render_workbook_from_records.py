from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def apply_table_style(table: Table) -> None:
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )


def reset_table(sheet, table_name: str, header_count: int, row_count: int) -> None:
    try:
        table_names = [table.displayName for table in sheet.tables.values()]
        for existing_name in table_names:
            del sheet.tables[existing_name]
    except Exception:
        sheet.tables.clear()

    end_row = max(4, 3 + row_count)
    end_col = get_column_letter(header_count)
    table = Table(displayName=table_name, ref=f"A3:{end_col}{end_row}")
    apply_table_style(table)
    sheet.add_table(table)


def clear_sheet_data(sheet) -> None:
    if sheet.max_row > 3:
        sheet.delete_rows(4, sheet.max_row - 3)


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("Usage: render_workbook_from_records.py <template_path> <records_json_path> <output_path>")

    template_path = Path(sys.argv[1])
    records_path = Path(sys.argv[2])
    output_path = Path(sys.argv[3])

    workbook = load_workbook(template_path)
    sheet_exports = json.loads(records_path.read_text(encoding="utf8"))

    for sheet_export in sheet_exports:
        sheet = workbook[sheet_export["sheetName"]]
        clear_sheet_data(sheet)

        headers = sheet_export["headers"]
        rows = sheet_export["rows"] or [[]]

        for column_index, header in enumerate(headers, start=1):
            sheet.cell(row=3, column=column_index, value=header)

        if not sheet_export["rows"]:
            rows = [["" for _ in headers]]

        for row_index, row_values in enumerate(rows, start=4):
            for column_index, value in enumerate(row_values, start=1):
                sheet.cell(row=row_index, column=column_index, value=value)

        if sheet_export.get("tableName"):
            reset_table(sheet, sheet_export["tableName"], len(headers), len(rows))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


if __name__ == "__main__":
    main()
